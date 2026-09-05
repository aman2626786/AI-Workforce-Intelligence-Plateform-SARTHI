import pymupdf as fitz
import docx
from pathlib import Path

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_resumes"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

def generate_standard_pdf():
    doc = fitz.open()
    page = doc.new_page()
    
    text = """Yogesh Kumar
Email: yogesh.kumar@example.com | Phone: +91 9876543210
Location: Jaipur, Rajasthan | LinkedIn: linkedin.com/in/yogesh-kumar | GitHub: github.com/yogesh-k

EDUCATION
B.Tech in Computer Science and Engineering
National Institute of Technology
Graduation: 2026
CGPA: 8.6 / 10

SKILLS
Technical Skills: Python, SQL, PostgreSQL, Pandas, NumPy, Power BI, Scikit-Learn, Machine Learning, Docker, Git

EXPERIENCE
Data Analyst Intern | TechCorp Analytics
May 2025 - July 2025
- Built SQL queries and automated business intelligence dashboards.
- Cleaned and manipulated 100k+ customer records using Python and Pandas.

PROJECTS
Customer Churn Prediction Model
- Developed an ML pipeline using Scikit-Learn, XGBoost, and Python.
- Achieved 89% accuracy in predicting subscriber churn.
URL: https://github.com/yogesh-k/churn-prediction

Interactive Sales Intelligence Dashboard
- Modeled corporate sales dataset using Power BI and DAX measures.
- Built automated reporting pipelines using PostgreSQL.

CERTIFICATIONS
- Google Data Analytics Professional Certificate by Coursera
- SQL for Data Science Certificate by Coursera
"""
    # Insert text into page
    rect = fitz.Rect(50, 50, 550, 750)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv")
    
    out_path = SAMPLE_DIR / "standard_resume.pdf"
    doc.save(str(out_path))
    doc.close()
    print(f"Generated {out_path}")

def generate_alt_headings_docx():
    doc = docx.Document()
    
    doc.add_heading("Priya Sharma", level=1)
    doc.add_paragraph("Email: priya.sharma@example.com | Phone: 9812345678\nBengaluru, India | linkedin.com/in/priya-sharma | github.com/priya-dev")
    
    doc.add_heading("ACADEMIC QUALIFICATIONS", level=2)
    doc.add_paragraph("Bachelor of Technology in Information Technology\nArya College of Engineering\nYear of Passing: 2025 | CGPA: 8.2")
    
    doc.add_heading("TECHNICAL EXPERTISE", level=2)
    doc.add_paragraph("Python, React, TypeScript, Node.js, Express.js, MongoDB, Docker, Git, RESTful APIs")
    
    doc.add_heading("WORK HISTORY", level=2)
    doc.add_paragraph("Software Engineering Intern at CloudWorks Systems (Jan 2025 - Jun 2025)\n- Developed full-stack web applications using React, TypeScript, and Node.js.")
    
    doc.add_heading("KEY PROJECTS", level=2)
    doc.add_paragraph("Real-Time Collaborative Code Editor\n- Built web application with Next.js, Express, and Docker containerization.")
    
    doc.add_heading("COURSES & CERTIFICATIONS", level=2)
    doc.add_paragraph("- Full Stack Web Development by Udemy\n- AWS Certified Cloud Practitioner by AWS")
    
    out_path = SAMPLE_DIR / "alt_headings_resume.docx"
    doc.save(str(out_path))
    print(f"Generated {out_path}")

def generate_conflicting_resume_pdf():
    doc = fitz.open()
    page = doc.new_page()
    
    # Note: Resume has city = "Bengaluru", while student input in registration will be "Jaipur"
    text = """Amit Verma
Email: amit.verma@example.com | Phone: +91 9123456780
Current Location: Bengaluru, Karnataka
LinkedIn: linkedin.com/in/amit-verma-ai | GitHub: github.com/amit-v

EDUCATION
B.Tech in Artificial Intelligence & Data Science
Manipal University
Graduation: 2025
CGPA: 7.9

TECHNICAL SKILLS
Python, SQL, PyTorch, Deep Learning, Generative AI, LangChain, BigQuery, Fastapi

PROJECTS
LLM-Powered Data Assistant
- Built generative AI query synthesizer using LangChain and Python.
- Automated SQL generation from natural language questions.

CERTIFICATIONS
- Deep Learning Specialization by Coursera
"""
    rect = fitz.Rect(50, 50, 550, 750)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv")
    
    out_path = SAMPLE_DIR / "conflicting_resume.pdf"
    doc.save(str(out_path))
    doc.close()
    print(f"Generated {out_path}")

if __name__ == "__main__":
    generate_standard_pdf()
    generate_alt_headings_docx()
    generate_conflicting_resume_pdf()
